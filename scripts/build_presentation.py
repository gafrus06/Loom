from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt


TITLE_SLIDE_TEXT = {
    "student_label": "Обучающийся:  Гафиятов Руслан Жамильевич",
    "supervisor_label": "Руководитель:  Воронкин Евгений Юрьевич",
    "topic": (
        "Разработка веб-приложения для организации отрядной деятельности\n"
        "и методической поддержки вожатых в детском оздоровительном лагере"
    ),
}


CONTENT_SLIDES = [
    {
        "title": "Актуальность",
        "body": [
            "Проблема:",
            "Детские лагеря до сих пор используют разрозненные инструменты: соцсети для новостей, бумажные методички, отдельные чаты и таблицы.",
            "Это затрудняет координацию вожатых, ведение отрядов и быстрый доступ к нужной информации во время смены.",
            "Статистика:",
            "1 октября 2024 года сообщалось примерно о 39 тыс. лагерей всех типов и 5,7 млн детей.",
            "17 апреля 2025 года сообщалось, что более 40,5 тыс. лагерей должны принять около 6 млн детей.",
            "Предлагаемое решение:",
            "Создание единой цифровой среды лагеря, объединяющей организацию работы, методическую поддержку и внутреннюю коммуникацию.",
        ],
    },
    {
        "title": "Цель и задачи",
        "body": [
            "Цель:",
            "Разработать веб-приложение, объединяющее организацию отрядной деятельности, коммуникацию сотрудников и методическую поддержку вожатых.",
            "Задачи:",
            "Проанализировать проблемы цифровизации детских лагерей.",
            "Изучить существующие аналоги и определить их ограничения.",
            "Спроектировать роли пользователей и контекст работы: лагерь, смена, отряд.",
            "Реализовать ключевые модули системы и обосновать выбранный стек технологий.",
        ],
    },
    {
        "title": "Анализ предметной области",
        "body": [
            "В предметной области можно выделить три группы решений:",
            "Сайты и страницы лагерей в соцсетях, ориентированные в основном на публикацию новостей.",
            "Зарубежные camp-management платформы, ориентированные на учет, регистрацию и работу с родителями.",
            "Локальные методические материалы в PDF, бумажных папках и чатах.",
            "Вывод:",
            "Большинство решений закрывает отдельные задачи, но не объединяет методическую поддержку вожатых и повседневную работу с отрядом.",
        ],
    },
    {
        "title": "Аналоги и сравнение",
        "body": [
            "Примеры аналогов:",
            "Campminder, CIRCUITREE, CampOrganizer.",
            "Сильные стороны аналогов:",
            "Регистрация участников, parent portal, административный учет, коммуникации.",
            "Ограничения аналогов для моей задачи:",
            "Нет фокуса на отрядной деятельности и оперативной методической поддержке вожатого.",
            "Преимущество моего решения:",
            "Система строится вокруг связки лагерь - смена - отряд и объединяет администрацию, старших вожатых, вожатых, медработников и родителей.",
        ],
    },
    {
        "title": "Техническое задание",
        "body": [
            "Основные требования к системе:",
            "Поддержка лагерей, смен, отрядов и карточек детей в едином контуре.",
            "Разграничение прав доступа по ролям и контексту пользователя.",
            "Наличие новостной ленты, задач смены, календаря и журнала отряда.",
            "Интеграция методических материалов для вожатых с быстрым доступом во время смены.",
            "Поддержка внутренней коммуникации без опоры только на внешние платформы.",
        ],
    },
    {
        "title": "Стек технологий",
        "body": [
            "Frontend: React.",
            "Backend: Java 21 и Spring Boot.",
            "Архитектура: микросервисная.",
            "Инфраструктурные компоненты: API Gateway и Eureka.",
            "Хранение данных: PostgreSQL.",
            "Кеширование и ускорение ленты: Redis.",
            "Событийное взаимодействие: Kafka.",
            "Контейнеризация и локальный запуск: Docker.",
        ],
    },
    {
        "title": "Результаты проектирования",
        "body": [
            "Ключевая идея проектирования:",
            "Система учитывает не только глобальную роль пользователя, но и его рабочий контекст.",
            "Контекст включает три уровня:",
            "Лагерь.",
            "Смена.",
            "Отряд.",
            "За счет этого пользователь получает только релевантные данные и функции для своей текущей работы.",
            "Рекомендуемая иллюстрация: схема архитектуры и схема контекста доступа.",
        ],
    },
    {
        "title": "Результат разработки",
        "body": [
            "В системе реализованы основные модули:",
            "Управление лагерями, сменами и отрядами.",
            "Карточки детей и состав отряда.",
            "Журнал отряда и задачи смены.",
            "Календарь событий.",
            "Внутренняя новостная лента лагеря.",
            "Рекомендуемые скриншоты: главная страница, страница отряда, панель старшего вожатого.",
        ],
    },
    {
        "title": "Разработка плана тестирования ",
        "body": [
            "Для проверки системы целесообразно использовать:",
            "Ручное функциональное тестирование пользовательских сценариев.",
            "Проверку ролевой модели и доступа к данным в разных контекстах.",
            "Тестирование создания и модерации постов.",
            "Проверку сценариев работы с методическими материалами.",
            "Проверку устойчивости взаимодействия между сервисами.",
        ],
    },
    {
        "title": "Результаты тестирования",
        "body": [
            "Подтверждаются следующие возможности системы:",
            "Корректная работа ролей администратора, вожатого, старшего вожатого, медработника и родителя.",
            "Отображение данных в зависимости от лагеря, смены и отряда.",
            "Создание, редактирование, закрепление и модерация постов.",
            "Работа методической панели: игры, упражнения, огоньки, возрастные рекомендации.",
            "Стабильная работа основных модулей в рамках микросервисной архитектуры.",
        ],
    },
    {
        "title": "Практическая значимость",
        "body": [
            "Практическая значимость проекта:",
            "Разработанное приложение может стать основой для цифровизации лагеря как единой рабочей среды.",
            "Ожидаемый эффект:",
            "Снижение количества разрозненных инструментов.",
            "Упрощение координации сотрудников.",
            "Ускорение доступа вожатых к методическим материалам.",
            "Повышение качества внутренней коммуникации лагеря.",
        ],
    },
    {
        "title": "Возможное развитие темы",
        "body": [
            "Направления дальнейшего развития:",
            "Расширение родительского кабинета.",
            "Добавление мобильной версии приложения.",
            "Развитие аналитики по сменам и отрядам.",
            "Интеллектуальные рекомендации для вожатых на основе контекста отряда.",
            "Интеграция с внешними сервисами оповещений и медиаконтента.",
        ],
    },
    {
        "title": "Заключение",
        "body": [
            "Цель работы достигнута.",
            "Разработано веб-приложение для организации отрядной деятельности и методической поддержки вожатых в детском лагере.",
            "В ходе выполнения работы были решены следующие задачи:",
            "Проведен анализ проблем цифровизации лагерей и существующих аналогов.",
            "Спроектирована ролевая и контекстная модель системы.",
            "Реализованы основные модули: управление лагерем, отрядом, новостной лентой и методическими материалами.",
            "Обоснован стек технологий и показаны перспективы дальнейшего развития проекта.",
        ],
    },
]


def find_shape_by_text(slide, expected_text):
    for shape in slide.shapes:
        text = getattr(shape, "text", "").strip()
        if text == expected_text:
            return shape
    return None


def clear_text_frame(text_frame):
    while len(text_frame.paragraphs) > 1:
        p = text_frame.paragraphs[-1]
        p._element.getparent().remove(p._element)
    text_frame.paragraphs[0].clear()


def set_body_text(shape, lines):
    text_frame = shape.text_frame
    clear_text_frame(text_frame)
    text_frame.word_wrap = True

    for idx, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        for run in paragraph.runs:
            run.font.size = Pt(18)
            if line.endswith(":"):
                run.font.bold = True


def set_simple_text(shape, text, size=None, bold=None):
    text_frame = shape.text_frame
    clear_text_frame(text_frame)
    paragraph = text_frame.paragraphs[0]
    paragraph.text = text
    for run in paragraph.runs:
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold


def set_title_text(slide, text):
    title_shape = None
    min_top = None
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        current_text = getattr(shape, "text", "")
        if not current_text:
            continue
        if current_text.strip().isdigit():
            continue
        if min_top is None or shape.top < min_top:
            title_shape = shape
            min_top = shape.top
    if title_shape is not None:
        set_simple_text(title_shape, text, size=24, bold=True)


def set_main_body(slide, lines):
    candidates = []
    title_shape = None
    min_top = None

    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = getattr(shape, "text", "").strip()
        if not text or text.isdigit():
            continue
        if min_top is None or shape.top < min_top:
            title_shape = shape
            min_top = shape.top

    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        if shape == title_shape:
            continue
        text = getattr(shape, "text", "").strip()
        if text.isdigit():
            continue
        candidates.append(shape)

    if not candidates:
        return

    for shape in candidates:
        clear_text_frame(shape.text_frame)

    body_shape = max(candidates, key=lambda item: item.width * item.height)
    set_body_text(body_shape, lines)


def set_title_slide(slide):
    theme_shape = find_shape_by_text(slide, "Тема")
    if theme_shape:
        set_simple_text(theme_shape, TITLE_SLIDE_TEXT["topic"], size=24, bold=True)

    person_shape = None
    for shape in slide.shapes:
        text = getattr(shape, "text", "")
        if "Обучающийся:" in text and "Руководитель:" in text:
            person_shape = shape
            break
    if person_shape:
        set_simple_text(
            person_shape,
            f"{TITLE_SLIDE_TEXT['student_label']}\n\n{TITLE_SLIDE_TEXT['supervisor_label']}",
            size=18,
            bold=False,
        )


def set_thank_you_slide(slide):
    theme_shape = find_shape_by_text(slide, "Тема")
    if theme_shape:
        set_simple_text(theme_shape, "Спасибо за внимание", size=28, bold=True)

    person_shape = None
    for shape in slide.shapes:
        text = getattr(shape, "text", "")
        if "Обучающийся:" in text and "Руководитель:" in text:
            person_shape = shape
            break
    if person_shape:
        set_simple_text(
            person_shape,
            "Обучающийся:  Гафиятов Руслан Жамильевич\n\nРуководитель:  Воронкин Евгений Юрьевич",
            size=18,
            bold=False,
        )


def remove_extra_slides(prs, keep_count):
    while len(prs.slides) > keep_count:
        slide_id_list = prs.slides._sldIdLst
        slide_id = slide_id_list[-1]
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        slide_id_list.remove(slide_id)


def main():
    base_dir = Path(r"C:/Users/gafru/Downloads/Telegram Desktop")
    template_path = base_dir / "ПРИМЕР ВКР БАКАЛАВР.pptx"
    output_path = base_dir / "Презентация_Гафиятов_ВКР.pptx"

    prs = Presentation(str(template_path))

    required_slides = len(CONTENT_SLIDES) + 2
    remove_extra_slides(prs, required_slides)

    set_title_slide(prs.slides[0])

    for idx, slide_data in enumerate(CONTENT_SLIDES, start=1):
        slide = prs.slides[idx]
        set_title_text(slide, slide_data["title"])
        set_main_body(slide, slide_data["body"])

    set_thank_you_slide(prs.slides[-1])
    prs.save(str(output_path))
    print(output_path)


if __name__ == "__main__":
    main()
